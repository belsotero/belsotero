"""Geração automática de relatórios técnicos e planilhas de inventário.

Este módulo combina boas práticas de padronização com automação para
agilizar a criação de relatórios patrimoniais completos, com texto,
tabelas e imagens anotadas.
"""

from __future__ import annotations

import argparse
import json
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Iterable, List, Mapping, MutableMapping, Sequence

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image, ImageDraw


@dataclass
class AnnotationColor:
    """Define uma paleta consistente para anotações."""

    box: str = "#ff6b6b"
    arrow: str = "#4dabf7"
    text: str = "#1c7ed6"


DEFAULT_COLORS = AnnotationColor()


def load_asset_data(path: Path) -> MutableMapping[str, object]:
    """Carrega os dados do ativo a partir de um arquivo JSON."""

    with path.open("r", encoding="utf-8") as file:
        return json.load(file)


def _ensure_output_dir(path: Path) -> None:
    path.mkdir(parents=True, exist_ok=True)


def _draw_arrow(draw: ImageDraw.ImageDraw, start: Sequence[int], end: Sequence[int], color: str, width: int = 5) -> None:
    """Desenha uma seta simples ligando dois pontos."""

    start_xy = (int(start[0]), int(start[1]))
    end_xy = (int(end[0]), int(end[1]))
    draw.line([start_xy, end_xy], fill=color, width=width)

    # Cabeça da seta (triângulo)
    arrow_head_size = 10 + width
    dx = start_xy[0] - end_xy[0]
    dy = start_xy[1] - end_xy[1]
    length = max((dx ** 2 + dy ** 2) ** 0.5, 1)
    ux, uy = dx / length, dy / length
    left = (
        end_xy[0] + arrow_head_size * (-ux - uy),
        end_xy[1] + arrow_head_size * (-uy + ux),
    )
    right = (
        end_xy[0] + arrow_head_size * (-ux + uy),
        end_xy[1] + arrow_head_size * (-uy - ux),
    )
    draw.polygon([end_xy, left, right], fill=color)


def annotate_image(
    image_path: Path,
    output_dir: Path,
    annotations: Iterable[Mapping[str, object]],
    colors: AnnotationColor = DEFAULT_COLORS,
) -> Path:
    """Cria uma cópia anotada da imagem com setas, caixas e textos."""

    _ensure_output_dir(output_dir)
    annotated_path = output_dir / f"{image_path.stem}_annotated{image_path.suffix}"

    with Image.open(image_path).convert("RGBA") as base_image:
        overlay = Image.new("RGBA", base_image.size, (255, 255, 255, 0))
        draw = ImageDraw.Draw(overlay)

        for annotation in annotations:
            kind = str(annotation.get("type", "")).lower()
            color = str(annotation.get("color") or getattr(colors, kind, colors.text))

            if kind == "arrow":
                _draw_arrow(draw, annotation["from"], annotation["to"], color)
            elif kind == "box":
                top_left = tuple(map(int, annotation["top_left"]))
                bottom_right = tuple(map(int, annotation["bottom_right"]))
                draw.rectangle([top_left, bottom_right], outline=color, width=int(annotation.get("width", 4)))
            elif kind == "text":
                position = tuple(map(int, annotation.get("position", (10, 10))))
                draw.text(position, str(annotation.get("label", "")), fill=color)

        composed = Image.alpha_composite(base_image, overlay)
        composed.convert("RGB").save(annotated_path, quality=95)

    return annotated_path


def _add_key_value_row(table, key: str, value: str) -> None:
    cells = table.add_row().cells
    cells[0].text = key
    cells[1].text = value
    for run in cells[0].paragraphs[0].runs:
        run.font.bold = True


def _format_table(table) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    table.style = "Light Shading Accent 1"


def _format_heading(paragraph, font_size: int = 18) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.runs[0]
    run.font.size = Pt(font_size)
    run.font.bold = True


def create_word_report(
    data: Mapping[str, object],
    output_path: Path,
    media_dir: Path,
    annotated_dir: Path,
    annotate: bool = True,
) -> Path:
    """Gera o documento principal do relatório técnico."""

    document = Document()
    asset = data.get("asset", {})

    title = document.add_heading(
        f"Relatório Técnico - {asset.get('name', 'Ativo sem nome')}", level=0
    )
    _format_heading(title)

    info_table = document.add_table(rows=0, cols=2)
    for label, key in (
        ("Código", "code"),
        ("Endereço", "address"),
        ("Data da inspeção", "inspection_date"),
        ("Responsável", "responsible"),
    ):
        value = str(asset.get(key, ""))
        if key == "inspection_date" and value:
            try:
                value = datetime.fromisoformat(value).strftime("%d/%m/%Y")
            except ValueError:
                pass
        _add_key_value_row(info_table, label, value)
    _format_table(info_table)

    document.add_paragraph()

    summary = data.get("summary", [])
    if summary:
        summary_heading = document.add_heading("Resumo Executivo", level=1)
        _format_heading(summary_heading, font_size=16)
        for item in summary:
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.add_run(f"{item.get('topic', 'Item')}: ").bold = True
            paragraph.add_run(str(item.get("notes", "")))

    inspections: List[Mapping[str, object]] = list(data.get("inspections", []))
    for section in inspections:
        heading = document.add_heading(str(section.get("location", "Ambiente")), level=1)
        _format_heading(heading, font_size=16)

        table = document.add_table(rows=1, cols=4)
        headers = ["Elemento", "Condição", "Prioridade", "Recomendação"]
        for cell, text in zip(table.rows[0].cells, headers):
            cell.text = text
            run = cell.paragraphs[0].runs[0]
            run.font.bold = True
        _format_table(table)

        for element in section.get("elements", []):
            cells = table.add_row().cells
            cells[0].text = str(element.get("element", ""))
            cells[1].text = str(element.get("status", ""))
            cells[2].text = str(element.get("priority", ""))
            cells[3].text = str(element.get("recommendation", ""))

    photos: List[Mapping[str, object]] = list(data.get("photos", []))
    if photos:
        photos_heading = document.add_heading("Registro Fotográfico", level=1)
        _format_heading(photos_heading, font_size=16)

        for photo in photos:
            original_path = media_dir / str(photo.get("file", ""))
            if not original_path.exists():
                document.add_paragraph(f"Imagem não localizada: {original_path}")
                continue

            annotations = photo.get("annotations", []) if annotate else []
            image_to_insert = (
                annotate_image(original_path, annotated_dir, annotations)
                if annotations
                else original_path
            )

            paragraph = document.add_paragraph()
            run = paragraph.add_run()
            run.add_picture(str(image_to_insert), width=Inches(5.5))
            caption = document.add_paragraph(str(photo.get("caption", "")))
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption.style = "Caption"

    document.save(output_path)
    return output_path


def create_excel_inventory(data: Mapping[str, object], output_path: Path) -> Path:
    """Gera uma planilha de acompanhamento com resumo das inspeções."""

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Inventário"

    headers = ["Local", "Elemento", "Condição", "Prioridade", "Recomendação", "Foto"]
    sheet.append(headers)

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=column)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = alignment
        sheet.column_dimensions[get_column_letter(column)].width = 30

    for section in data.get("inspections", []):
        location = section.get("location", "")
        for element in section.get("elements", []):
            sheet.append(
                [
                    location,
                    element.get("element", ""),
                    element.get("status", ""),
                    element.get("priority", ""),
                    element.get("recommendation", ""),
                    element.get("photo", ""),
                ]
            )

    workbook.save(output_path)
    return output_path


def generate_reports(
    data_path: Path,
    media_dir: Path,
    output_dir: Path,
    annotate: bool = True,
) -> Mapping[str, Path]:
    """Fluxo completo para gerar documento Word e planilha Excel padronizados."""

    data = load_asset_data(data_path)
    _ensure_output_dir(output_dir)
    annotated_dir = output_dir / "annotated"
    _ensure_output_dir(annotated_dir)

    word_path = output_dir / "relatorio_tecnico.docx"
    excel_path = output_dir / "inventario_inspecao.xlsx"

    create_word_report(data, word_path, media_dir, annotated_dir, annotate=annotate)
    create_excel_inventory(data, excel_path)

    return {"word": word_path, "excel": excel_path}


def _build_argument_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Gera relatórios técnicos padronizados.")
    parser.add_argument("data", type=Path, help="Arquivo JSON com os dados do ativo.")
    parser.add_argument(
        "--media-dir",
        type=Path,
        default=Path("media"),
        help="Pasta onde as imagens originais estão armazenadas.",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=Path("output"),
        help="Pasta onde os arquivos gerados serão salvos.",
    )
    parser.add_argument(
        "--sem-anotacoes",
        action="store_true",
        help="Desativa a criação de imagens anotadas.",
    )
    return parser


def main(args: Sequence[str] | None = None) -> Mapping[str, Path]:
    parser = _build_argument_parser()
    parsed = parser.parse_args(args=args)

    outputs = generate_reports(
        data_path=parsed.data,
        media_dir=parsed.media_dir,
        output_dir=parsed.output_dir,
        annotate=not parsed.sem_anotacoes,
    )

    print("Arquivos gerados:")
    for label, path in outputs.items():
        print(f" - {label}: {path}")
    return outputs


if __name__ == "__main__":
    main()
